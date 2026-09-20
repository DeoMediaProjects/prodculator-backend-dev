"""Every recorded claim the engines cannot act on, with what it needs to become.

Usage::

    DB_URL=… venv/Scripts/python scripts/build_correction_workbook.py --out fixes.xlsx
    DB_URL=… venv/Scripts/python scripts/build_correction_workbook.py --apply fixes.xlsx
    DB_URL=… venv/Scripts/python scripts/build_correction_workbook.py --apply fixes.xlsx --commit

WHY THIS IS NOT A GUESS
-----------------------
Most of these claims fail for one reason: a rule names a section its festival's
deadlines never mentioned. "Competition" against "Feature Competition",
"Feature" against "Feature-length films". The two are almost certainly the same
section, and "almost certainly" is exactly the standard that turned "BC Arts
Council" into "California Arts Council" when this project last let a similarity
score decide something.

So this matches nothing. It puts the sections that have a dated deadline in a
column beside the sections the rules name, and leaves the correction blank. The
person who wrote both lists settles it in seconds; nobody else should try.

WHAT A CORRECTION COSTS
-----------------------
A corrected claim goes back to PENDING and loses its reviewer and its QA
signature. That is not friction for its own sake: the people who signed the old
reading did not sign the new words, and carrying their names forward would put a
signature on something they never saw. For festival section rules and market
hard gates that means a second reviewer again.

The previous reading is written into the claim's notes before the new one lands,
so the fact that someone once read the source differently survives the fix.
"""

from __future__ import annotations

import argparse
import os
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import sqlalchemy as sa
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

CONTEXT_COLUMNS = (
    "gate",
    "subject_id",
    "subject_name",
    "field",
    "review_state",
    "lines_that_produce_nothing",
    "sections_with_a_dated_deadline",
    "current_value",
)
CORRECTION_COLUMNS = (
    "corrected_value",
    "verified_by",
    "verified_on",
    "source_url",
    "source_basis",
)

_HEADER = PatternFill("solid", fgColor="1F3B57")
_FILL = PatternFill("solid", fgColor="FFF3CD")

INSTRUCTIONS = [
    ("Corrections", True),
    ("", False),
    ("Each row is one recorded claim that the engines cannot act on. The claim is", False),
    ("not necessarily wrong — most of these were read correctly off an official", False),
    ("page and then written in a shape no parser can follow.", False),
    ("", False),
    ("Read two columns first:", False),
    ("  lines_that_produce_nothing     which lines failed, and why.", False),
    ("  sections_with_a_dated_deadline the section names this festival's deadlines", False),
    ("                                 actually use. A rule has to name one of these", False),
    ("                                 exactly, or it attaches to nothing.", False),
    ("", False),
    ("Then rewrite the WHOLE value in corrected_value — every line, not just the", False),
    ("broken ones. The corrected value replaces the recorded one entirely.", False),
    ("", False),
    ("Fill verified_by, verified_on, source_url and source_basis as well. A", False),
    ("correction is a new assertion about the source and carries its own evidence.", False),
    ("", False),
    ("A corrected claim goes back to PENDING and loses its reviewer and QA", False),
    ("signature. The people who signed the old reading did not sign the new words.", False),
    ("For festival section rules and market hard gates that means a second", False),
    ("reviewer again.", False),
    ("", False),
    ("Leave a row blank to leave the claim alone.", False),
    ("", False),
    ("Apply with:", False),
    ("    DB_URL=... venv/Scripts/python scripts/build_correction_workbook.py \\", False),
    ("        --apply THIS.xlsx --commit", False),
]


def _subject_names(engine: sa.Engine) -> dict[str, str]:
    """Reader-facing names for festivals, market tracks and companies."""
    names: dict[str, str] = {}
    with engine.connect() as conn:
        present = set(sa.inspect(conn).get_table_names())
        if "engine_handoff_records" in present:
            for row in conn.execute(
                sa.text("SELECT record_id, name FROM engine_handoff_records")
            ):
                names.setdefault(str(row[0]), str(row[1]))
        if "market_tracks" in present:
            for row in conn.execute(sa.text("SELECT id, name FROM market_tracks")):
                names.setdefault(str(row[0]), str(row[1]))
        if "incentive_programs" in present:
            for row in conn.execute(sa.text("SELECT id, program FROM incentive_programs")):
                names.setdefault(str(row[0]), str(row[1] or ""))
    return names


def _dated_sections(claims, *, today: date) -> dict[str, list[str]]:
    """Per festival, the section names that carry a deadline the engine read."""
    from app.modules.reports.verification_ledger import SourceClaim
    from app.modules.reports.verified_cycles import build_cycles

    as_if = [
        SourceClaim(**{
            **claim.__dict__,
            "review_state": "VERIFIED",
            "reviewed_by": "preview",
            "qa_by": "preview-qa" if claim.needs_independent_qa else claim.qa_by,
        })
        for claim in claims
    ]
    cycles, _ = build_cycles(as_if, today=today)
    sections: dict[str, list[str]] = defaultdict(list)
    for cycle in cycles:
        if cycle.section_name:
            sections[cycle.record_id].append(cycle.section_name)
    return {key: sorted(value) for key, value in sections.items()}


def _problems_by_claim(claims, *, today: date) -> dict[tuple[str, str, str], list[str]]:
    from app.modules.reports.verification_ledger import SourceClaim
    from app.modules.reports.verified_cycles import build_cycles

    as_if = [
        SourceClaim(**{
            **claim.__dict__,
            "review_state": "VERIFIED",
            "reviewed_by": "preview",
            "qa_by": "preview-qa" if claim.needs_independent_qa else claim.qa_by,
        })
        for claim in claims
    ]
    _, problems = build_cycles(as_if, today=today)
    # Keyed by the claim the line came from, not by its subject. A festival has
    # two claims and one of them is usually blameless; sending both back
    # because the other has a bad line is how correct work gets rewritten.
    found: dict[tuple[str, str, str], list[str]] = defaultdict(list)
    for problem in problems:
        if not problem.needs_correction:
            continue
        key = (problem.gate, problem.subject_id, problem.field)
        found[key].append(f"{problem.detail[:70]!r} — {problem.reason}")
    return found


def needs_correction(engine: sa.Engine, *, today: date) -> list[tuple]:
    """Every claim with at least one line the engines cannot act on."""
    from app.modules.reports import verification_store as store
    from app.modules.reports.verification_ledger import REJECTED

    claims = [c for c in store.load_claims(engine) if c.review_state != REJECTED]
    problems = _problems_by_claim(claims, today=today)
    sections = _dated_sections(claims, today=today)
    names = _subject_names(engine)

    rows = []
    for claim in claims:
        trouble = problems.get((claim.gate, claim.subject_id, claim.field))
        if not trouble:
            continue
        rows.append((
            claim, trouble,
            sections.get(claim.subject_id, []),
            names.get(claim.subject_id, ""),
        ))
    return rows


def build(rows) -> Workbook:
    book = Workbook()
    sheet = book.active
    sheet.title = "How to fix"
    for index, (text, bold) in enumerate(INSTRUCTIONS, start=1):
        cell = sheet.cell(row=index, column=1, value=text)
        cell.font = Font(bold=bold or text.endswith(":"), size=13 if bold else 11)
    sheet.column_dimensions["A"].width = 96

    by_gate: dict[str, list] = defaultdict(list)
    for row in rows:
        by_gate[row[0].gate].append(row)

    headers = [*CONTEXT_COLUMNS, *CORRECTION_COLUMNS]
    for gate in sorted(by_gate):
        page = book.create_sheet(gate[:31])
        page.cell(row=1, column=1, value=f"{gate} — {len(by_gate[gate])} claim(s) to correct")
        page.cell(row=1, column=1).font = Font(bold=True, size=12)
        for column, name in enumerate(headers, start=1):
            cell = page.cell(row=2, column=column, value=name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = _HEADER
            cell.alignment = Alignment(vertical="center", wrap_text=True)

        entries = sorted(by_gate[gate], key=lambda item: (item[3], item[0].field))
        for offset, (claim, trouble, sections, name) in enumerate(entries):
            line = 3 + offset
            values = (
                claim.gate,
                claim.subject_id,
                name,
                claim.field,
                claim.review_state,
                "\n".join(trouble),
                "\n".join(sections) or "(none carry a date)",
                str(claim.value or ""),
            )
            for column, value in enumerate(values, start=1):
                cell = page.cell(row=line, column=column, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            for column in range(len(CONTEXT_COLUMNS) + 1, len(headers) + 1):
                page.cell(row=line, column=column).fill = _FILL

        widths = {
            "subject_id": 38, "subject_name": 34, "lines_that_produce_nothing": 60,
            "sections_with_a_dated_deadline": 30, "current_value": 52,
            "corrected_value": 52, "source_url": 34, "source_basis": 34,
        }
        for column, name in enumerate(headers, start=1):
            page.column_dimensions[get_column_letter(column)].width = widths.get(name, 16)
        page.freeze_panes = "A3"
    return book


def read_corrections(path: Path):
    from app.modules.reports.verification_ledger import SourceClaim

    book = load_workbook(path, data_only=True)
    claims, unreadable = [], []
    for name in book.sheetnames:
        page = book[name]
        header = [c.value for c in page[2]] if page.max_row >= 2 else []
        if "gate" not in (header or []) or "corrected_value" not in (header or []):
            continue
        index = {h: n for n, h in enumerate(header)}
        for number, raw in enumerate(page.iter_rows(min_row=3, values_only=True), start=3):
            corrected = raw[index["corrected_value"]]
            if corrected is None or not str(corrected).strip():
                continue
            verified_on = raw[index["verified_on"]]
            if isinstance(verified_on, datetime):
                verified_on = verified_on.date()
            elif not isinstance(verified_on, date):
                try:
                    verified_on = date.fromisoformat(str(verified_on or "")[:10])
                except ValueError:
                    unreadable.append((name, number, "verified_on is missing or not a date"))
                    continue
            claims.append(SourceClaim(
                gate=str(raw[index["gate"]] or "").strip(),
                subject_id=str(raw[index["subject_id"]] or "").strip(),
                field=str(raw[index["field"]] or "").strip(),
                value=str(corrected).strip(),
                source_url=str(raw[index["source_url"]] or "").strip(),
                source_basis=str(raw[index["source_basis"]] or "").strip(),
                verified_on=verified_on,
                verified_by=str(raw[index["verified_by"]] or "").strip(),
            ))
    return claims, unreadable


def _engine():
    db_url = os.environ.get("DB_URL")
    if not db_url:
        from app.core.config import get_settings

        db_url = get_settings().DB_URL
    return sa.create_engine(db_url)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, help="Where to write the correction workbook")
    parser.add_argument("--apply", type=Path, help="Read corrections from a filled workbook")
    parser.add_argument("--commit", action="store_true", help="With --apply, write them")
    parser.add_argument("--show", type=int, default=20)
    args = parser.parse_args()
    today = date.today()

    if args.apply:
        from app.modules.reports import verification_store as store

        claims, unreadable = read_corrections(args.apply)
        result = store.supersede_claims(
            _engine(), claims, apply=args.commit, today=today
        )
        print(f"Corrections read : {len(claims)}")
        print(f"  {'replaced' if args.commit else 'would replace'} : {len(result.replaced)}")
        print(f"  identical to what is recorded : {result.unchanged}")
        print(f"  refused                       : {len(result.refused)}")
        print(f"  naming no recorded claim      : {len(result.missing)}")
        print(f"  unreadable rows               : {len(unreadable)}")
        for claim, problems in result.refused[: args.show]:
            print(f"  {claim.gate}/{claim.subject_id}/{claim.field}")
            for problem in problems:
                print(f"      - {problem}")
        for claim in result.missing[: args.show]:
            print(f"  no recorded claim for {claim.gate}/{claim.subject_id}/{claim.field}")
        for sheet, number, why in unreadable[: args.show]:
            print(f"  {sheet} row {number}: {why}")
        print()
        if args.commit:
            print(
                "Replaced claims are PENDING again and carry no reviewer. "
                "They need reviewing before any engine reads them."
            )
        else:
            print("Preflight only. Re-run with --commit to replace these readings.")
        raise SystemExit(0 if result.is_clean and not unreadable else 1)

    if not args.out:
        parser.exit(2, "Give --out for a new correction workbook, or --apply a filled one\n")

    engine = _engine()
    rows = needs_correction(engine, today=today)
    if not rows:
        parser.exit(0, "Nothing recorded needs correcting.\n")

    build(rows).save(args.out)
    print(f"Correction workbook: {args.out}")
    print(f"  claims needing a rewrite : {len(rows)}")
    print()
    by_gate: dict[str, int] = defaultdict(int)
    for claim, _, _, _ in rows:
        by_gate[f"{claim.gate} / {claim.field}"] += 1
    for label, count in sorted(by_gate.items()):
        print(f"  {label:44} {count:>4}")
    print()
    print("No section name is matched to another by this file. That is the point.")


if __name__ == "__main__":
    main()
