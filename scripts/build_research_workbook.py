"""Compile every outstanding verification gate into one research workbook.

Usage::

    DB_URL=… venv/Scripts/python scripts/build_research_workbook.py

Produces a single .xlsx: an instructions sheet, then one sheet per gate. The
CSVs from ``build_research_pack.py`` are the same data — this is the version a
person actually works in, which for six hundred rows of lookup means a
spreadsheet rather than a prose document.

WHY THE FILL-IN COLUMNS ARE SHADED
----------------------------------
Every sheet has context columns and five blank ones. The blanks are the work;
the context is what the system currently believes, which in the incentive gate
is precisely what is under review. Shading them is not decoration — a
researcher who types an answer into a ``ctx_`` column has silently overwritten
the thing the reviewer needed to compare against.
"""

from __future__ import annotations

import argparse
import os
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
CONTEXT_FILL = PatternFill("solid", fgColor="F2F2F2")
ANSWER_FILL = PatternFill("solid", fgColor="FFF2CC")
TITLE_FONT = Font(bold=True, size=16, color="1F3864")
H2_FONT = Font(bold=True, size=12, color="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

RESEARCH_COLUMNS = ("value", "source_url", "source_basis", "verified_on", "verified_by")


def _sheet(wb: Workbook, title: str, rows: list[dict], context: list[str], note: str):
    ws = wb.create_sheet(title[:31])
    ws.append([note])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3 + len(context) + 5)
    ws.cell(row=1, column=1).font = Font(italic=True, size=9, color="7F7F7F")
    ws.cell(row=1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 28

    header = ["gate", "subject_id", "field", *[f"ctx_{c}" for c in context], *RESEARCH_COLUMNS]
    ws.append(header)
    for index, _ in enumerate(header, start=1):
        cell = ws.cell(row=2, column=index)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    ws.row_dimensions[2].height = 30

    first_answer = 3 + len(context) + 1
    for row in rows:
        ws.append(
            [row["gate"], row["subject_id"], row["field"]]
            + [row.get(c, "") for c in context]
            + ["", "", "", "", ""]
        )

    for r in range(3, ws.max_row + 1):
        for c in range(1, len(header) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = ANSWER_FILL if c >= first_answer else CONTEXT_FILL

    widths = [14, 26, 20] + [30] * len(context) + [22, 38, 46, 13, 15]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    ws.freeze_panes = ws.cell(row=3, column=4)
    ws.auto_filter.ref = f"A2:{get_column_letter(len(header))}{max(ws.max_row, 2)}"
    return ws


_INSTRUCTIONS: list[tuple[str, str]] = [
    ("", "Seven gates stand between the v2 engines and a paid cutover. Each is a "
         "set of facts nobody has verified yet. One sheet per gate."),
    ("How to work a sheet",
     "Fill the five shaded columns on the right of each row: value, source_url, "
     "source_basis, verified_on, verified_by. Leave the grey ctx_ columns alone — "
     "they are what the system currently believes, and a reviewer compares your "
     "answer against them. In the incentive gate that belief is exactly what is "
     "under review."),
    ("Recording what you find",
     "Save the sheet as CSV, then:\n"
     "    python scripts/record_verifications.py <file>.csv          (preflight)\n"
     "    python scripts/record_verifications.py <file>.csv --apply\n"
     "Recording is not approving. Every claim lands PENDING and reaches an engine "
     "only once a named reviewer signs it off:\n"
     "    python scripts/record_verifications.py --review GATE/SUBJECT_ID/FIELD "
     "--reviewer <name> [--qa-by <other name>]"),
    ("A claim is rejected if",
     "• it has no source_url, or the URL is http rather than https\n"
     "• the URL points at our own database or admin — a record cannot verify itself\n"
     "• source_basis is blank\n"
     "• verified_on is a future date\n"
     "• for gates 5 and 7, the QA signature is the same person who made the claim"),
    ("What source_basis is for",
     "One sentence, in the source's own words, saying what the page actually "
     "states. It lets a reviewer check your reading without revisiting the page, "
     "and it makes a claim that has quietly drifted from its source visible."),
    ("Suggested order",
     "Gate 1 first: smallest, needs no second reviewer, and it is the only gate "
     "that makes the old-versus-v2 report comparison produce anything real — "
     "Section 07 stays empty until programmes can calculate.\n"
     "Then 2 and 3 (bounded, no second reviewer).\n"
     "Then 4 and 6 (transcription).\n"
     "Then 5 and 7 (interpretation, two people each)."),
]

_GATE_NOTES: dict[str, str] = {
    "1 Incentive engines": (
        "FIELD: qs_engine_type — which statutory formula the programme calculates on. "
        "ALLOWED: CORE_LOWER_OF (lower of local core and a % of global core) · "
        "ELIGIBLE_LOCAL_SPEND · QUALIFIED_LABOUR · QAPE · QNZPE · VFX_ONLY · PDV_ONLY · "
        "TIERED_SPEND · MULTI_BUCKET · INVESTOR_TAX_SHELTER · COMPETITIVE_GRANT · NO_PROGRAMME.    "
        "⚠ ctx_legacy_guess_qualifying_spend_type is the OLD column this work replaces. It is a "
        "guess, not an answer — copying it launders an old assumption into the new field and "
        "closes the gate without verifying anything. Read the statute. "
        "DONE WHEN the programme's official rules name the expenditure the rate applies to, and "
        "source_basis quotes that sentence."
    ),
    "2 Grants split parents": (
        "FIELD: successor_titles — which records in the 253-row v2 master replace this retired "
        "generic parent. Semicolon-separated exact titles. These were retired because the master "
        "carries the programme as several specific records; automatic matching could not find "
        "them and two similarity heuristics produced nonsense. "
        "⚠ Write NONE if the funding body genuinely has no record in the master. That is the "
        "important finding — it means a real programme is no longer reachable by producers."
    ),
    "3 Commercial profiles": (
        "FIELDS: primary_source_confirmation (CONFIRMED / CONTRADICTED — the profile came from a "
        "directory listing, check it against the company's own site) · acquisition_stage "
        "(development / pre_production / production / post_production / completed) · "
        "genre_specialties (semicolon-separated) · festival_market_signal (documented attendance).    "
        "⚠ A published contact page is NOT evidence that unsolicited submissions are accepted. If "
        "the site says nothing, that is UNKNOWN, not open."
    ),
    "4 Market cycles": (
        "FIELD: deadline — the current call's closing date, YYYY-MM-DD. Transcription from the "
        "official call page; no second reviewer needed. "
        "⚠ If the next call is not yet announced, write NOT_ANNOUNCED. Do not carry last year's "
        "date forward and do not infer one from a pattern."
    ),
    "5 Market hard gates": (
        "TWO PEOPLE. FIELD: hard_gates — the prose rules, typed. One per line: "
        "project_field operator expected. Operators: equals, one_of, at_least, at_most, "
        "greater_than, less_than, overlaps, manual_confirmation. "
        "Example: 'feature projects at development stage' becomes two lines — "
        "'format one_of feature' and 'stage one_of development'.    "
        "⚠ Use manual_confirmation whenever a rule needs human judgement. A rule typed as "
        "decidable when it is not is worse than leaving it unstructured, because the engine will "
        "then act on it. Leave anything you cannot type out entirely."
    ),
    "6 Festival deadlines": (
        "FIELD: section_deadlines — every section with its own deadline. One per line: "
        "'Section name | YYYY-MM-DD'. A festival has several sections closing on different dates "
        "— feature, short and episodic differ — which is why a festival-level date cannot answer "
        "when THIS production must submit. "
        "⚠ ctx_current_deadline_prose on 121 records literally reads 'Verify current submission "
        "windows on the…'. That is the source telling you it did not resolve them."
    ),
    "7 Festival rules": (
        "TWO PEOPLE. FIELD: section_rules — per section, eligibility typed as in gate 5, plus the "
        "premiere requirement as WORLD / INTERNATIONAL / NATIONAL / NONE. "
        "Example lines: 'Main Competition | premiere_requirement WORLD' and "
        "'Shorts | runtime_minutes at_most 40'.    "
        "⚠ NONE means the festival STATES it imposes no premiere requirement. A silent page is "
        "not NONE — leave it unrecorded, and the engine treats an unrecorded requirement as "
        "needing confirmation."
    ),
}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--out", type=Path, default=ROOT / "Prodculator_v2_Research_Pack.xlsx"
    )
    args = parser.parse_args()

    from scripts.build_research_pack import (  # noqa: PLC0415
        _CTX_COMMERCIAL,
        _CTX_FESTIVAL,
        _CTX_INCENTIVE,
        _CTX_MARKET,
        _CTX_SPLIT,
        commercial_rows,
        festival_rows,
        incentive_rows,
        market_rows,
        split_parent_rows,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Start here"
    ws.append(["Prodculator v2 — source verification research pack"])
    ws.cell(row=1, column=1).font = TITLE_FONT
    ws.append([])
    for heading, body in _INSTRUCTIONS:
        if heading:
            ws.append([heading])
            ws.cell(row=ws.max_row, column=1).font = H2_FONT
        ws.append([body])
        cell = ws.cell(row=ws.max_row, column=1)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[ws.max_row].height = 15 * (body.count("\n") + 3)
        ws.append([])
    ws.column_dimensions["A"].width = 130

    counts: dict[str, int] = {}
    db_url = os.environ.get("DB_URL")
    if db_url:
        import sqlalchemy as sa

        engine = sa.create_engine(db_url)
        with engine.connect() as conn:
            if conn.dialect.name == "postgresql":
                conn.execute(sa.text("SET TRANSACTION READ ONLY"))
            for title, rows, ctx in (
                ("1 Incentive engines", incentive_rows(conn), _CTX_INCENTIVE),
                ("2 Grants split parents", split_parent_rows(conn), _CTX_SPLIT),
            ):
                _sheet(wb, title, rows, ctx, _GATE_NOTES[title])
                counts[title] = len(rows)
    else:
        print("DB_URL not set — gates 1 and 2 need the live database; skipping them.")

    for title, rows, ctx in (
        ("3 Commercial profiles", commercial_rows(), _CTX_COMMERCIAL),
        ("4 Market cycles", market_rows("cycle"), _CTX_MARKET),
        ("5 Market hard gates", market_rows("gates"), _CTX_MARKET),
        ("6 Festival deadlines", festival_rows("deadline"), _CTX_FESTIVAL),
        ("7 Festival rules", festival_rows("rules"), _CTX_FESTIVAL),
    ):
        _sheet(wb, title, rows, ctx, _GATE_NOTES[title])
        counts[title] = len(rows)

    # Index on the instructions sheet, written last so the counts are real.
    ws.append(["Gates in this workbook"])
    ws.cell(row=ws.max_row, column=1).font = H2_FONT
    for title, count in counts.items():
        ws.append([f"    {title} — {count} rows"])
    ws.append([f"    TOTAL — {sum(counts.values())} rows"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)

    wb.save(args.out)
    print(f"Written: {args.out}")
    for title, count in counts.items():
        print(f"  {title:26} {count:>4} rows")
    print(f"  {'TOTAL':26} {sum(counts.values()):>4} rows")


if __name__ == "__main__":
    main()
